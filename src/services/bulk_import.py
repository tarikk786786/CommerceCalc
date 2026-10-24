"""
Bulk product import and calculation service for CommerceCalc.

Handles CSV/XLSX import, validation, batch calculation, and export.
"""
import csv
import io
from decimal import Decimal, InvalidOperation
from typing import List, Dict, Any, Optional, Tuple
from dataclasses import dataclass, field
from datetime import datetime
import openpyxl
from openpyxl import Workbook

from ..core.profit import ProductCost, SellingDetails, SellingExpenses, calculate_profit
from ..core.tax import GSTConfig
from ..core.money import to_decimal, money


# Required columns for bulk import
REQUIRED_COLUMNS = ["product_name", "purchase_cost", "actual_selling_price"]

# Full column specification
COLUMN_SPEC = {
    "product_name": {"type": "str", "required": True, "description": "Product name"},
    "sku": {"type": "str", "required": False, "description": "SKU/product code"},
    "purchase_cost": {"type": "decimal", "required": True, "description": "Product purchase cost (₹)"},
    "packaging_cost": {"type": "decimal", "required": False, "description": "Packaging cost (₹)"},
    "inbound_shipping": {"type": "decimal", "required": False, "description": "Inbound shipping cost (₹)"},
    "listed_price": {"type": "decimal", "required": False, "description": "Listed/MRP price (₹)"},
    "actual_selling_price": {"type": "decimal", "required": True, "description": "Actual selling price (₹)"},
    "quantity": {"type": "int", "required": False, "description": "Quantity to calculate"},
    "marketplace_commission": {"type": "decimal", "required": False, "description": "Marketplace commission (₹)"},
    "forward_shipping": {"type": "decimal", "required": False, "description": "Forward shipping cost (₹)"},
    "payment_gateway_fee": {"type": "decimal", "required": False, "description": "Payment gateway fee (₹)"},
    "advertising_cost": {"type": "decimal", "required": False, "description": "Advertising cost (₹)"},
    "gst_rate_pct": {"type": "decimal", "required": False, "description": "GST rate % (0, 5, 12, 18, 28)"},
    "category": {"type": "str", "required": False, "description": "Product category"},
    "channel": {"type": "str", "required": False, "description": "Sales channel"},
}


@dataclass
class ImportRow:
    """A single imported product row with validation status."""
    row_number: int
    raw_data: Dict[str, str] = field(default_factory=dict)
    parsed_data: Dict[str, Any] = field(default_factory=dict)
    errors: List[str] = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)
    is_valid: bool = True
    result: Any = None   # ProfitResult if calculated
    calculated: bool = False


@dataclass
class ImportBatchResult:
    """Result of a full batch import operation."""
    total_rows: int = 0
    valid_rows: int = 0
    failed_rows: int = 0
    calculated_rows: int = 0
    profitable_rows: int = 0
    loss_rows: int = 0
    rows: List[ImportRow] = field(default_factory=list)
    column_mapping: Dict[str, str] = field(default_factory=dict)
    timestamp: str = ""
    filename: str = ""
    errors: List[str] = field(default_factory=list)


def parse_decimal_safe(value: str, field_name: str = "") -> Tuple[Optional[Decimal], str]:
    """Parse decimal, returning (value, error_message)."""
    if not value or str(value).strip() == "":
        return Decimal("0"), ""
    try:
        d = to_decimal(str(value).strip().replace(",", "").replace("₹", "").replace("$", ""))
        if d < 0 and field_name in ["purchase_cost", "actual_selling_price"]:
            return None, f"{field_name}: negative values not allowed"
        return d, ""
    except (InvalidOperation, ValueError) as e:
        return None, f"{field_name}: invalid number '{value}'"


def validate_row(row_data: Dict[str, str], row_num: int) -> ImportRow:
    """Validate a single imported row."""
    import_row = ImportRow(row_number=row_num, raw_data=dict(row_data))

    # Check required columns
    for col in REQUIRED_COLUMNS:
        if col not in row_data or not str(row_data.get(col, "")).strip():
            import_row.errors.append(f"Missing required column: {col}")
            import_row.is_valid = False

    # Parse numeric fields
    parsed = {}
    for col, spec in COLUMN_SPEC.items():
        raw_val = row_data.get(col, "")
        if spec["type"] == "decimal":
            val, err = parse_decimal_safe(raw_val, col)
            if err:
                import_row.errors.append(err)
                import_row.is_valid = False
            else:
                parsed[col] = val or Decimal("0")
        elif spec["type"] == "int":
            try:
                parsed[col] = max(int(float(raw_val)) if raw_val else 1, 1)
            except (ValueError, TypeError):
                parsed[col] = 1
                import_row.warnings.append(f"{col}: invalid integer, defaulted to 1")
        else:
            parsed[col] = str(raw_val).strip()

    # Business validation
    if "actual_selling_price" in parsed and parsed["actual_selling_price"] <= 0:
        import_row.errors.append("actual_selling_price must be greater than 0")
        import_row.is_valid = False

    if "gst_rate_pct" in parsed:
        gst = parsed["gst_rate_pct"]
        if gst not in [Decimal("0"), Decimal("5"), Decimal("12"), Decimal("18"), Decimal("28")]:
            import_row.warnings.append(
                f"gst_rate_pct={gst} is not a standard preset rate. Verify this rate."
            )

    import_row.parsed_data = parsed
    return import_row


def calculate_batch(import_rows: List[ImportRow]) -> List[ImportRow]:
    """Calculate profitability for all valid import rows."""
    for row in import_rows:
        if not row.is_valid:
            continue
        try:
            parsed = row.parsed_data
            pc = ProductCost(
                purchase_cost=parsed.get("purchase_cost", Decimal("0")),
                packaging_cost=parsed.get("packaging_cost", Decimal("0")),
                inbound_shipping=parsed.get("inbound_shipping", Decimal("0")),
            )
            gst_config = GSTConfig(
                rate_pct=parsed.get("gst_rate_pct", Decimal("0")),
                inclusive=False,
            )
            sd = SellingDetails(
                listed_price=parsed.get("listed_price", parsed.get("actual_selling_price")),
                actual_selling_price=parsed.get("actual_selling_price"),
                quantity=int(parsed.get("quantity", 1)),
                gst_config=gst_config,
            )
            se = SellingExpenses(
                marketplace_commission=parsed.get("marketplace_commission", Decimal("0")),
                forward_shipping=parsed.get("forward_shipping", Decimal("0")),
                payment_gateway_fee=parsed.get("payment_gateway_fee", Decimal("0")),
                advertising_cost=parsed.get("advertising_cost", Decimal("0")),
            )
            row.result = calculate_profit(pc, sd, se)
            row.calculated = True
        except Exception as e:
            row.errors.append(f"Calculation error: {str(e)}")
            row.calculated = False
    return import_rows


def import_from_csv(file_path: str) -> ImportBatchResult:
    """Import and validate products from a CSV file."""
    batch = ImportBatchResult(timestamp=datetime.now().isoformat(), filename=file_path)

    try:
        with open(file_path, "r", encoding="utf-8-sig", newline="") as f:
            reader = csv.DictReader(f)
            headers = reader.fieldnames or []
            batch.column_mapping = {h.lower().strip().replace(" ", "_"): h for h in headers}

            rows_data = []
            for row in reader:
                # Normalize keys
                normalized = {k.lower().strip().replace(" ", "_"): v for k, v in row.items()}
                rows_data.append(normalized)

    except FileNotFoundError:
        batch.errors.append(f"File not found: {file_path}")
        return batch
    except Exception as e:
        batch.errors.append(f"Failed to read CSV: {str(e)}")
        return batch

    batch.total_rows = len(rows_data)
    import_rows = []
    for i, row_data in enumerate(rows_data, start=2):  # Row 1 is header
        import_row = validate_row(row_data, i)
        import_rows.append(import_row)

    # Calculate
    import_rows = calculate_batch(import_rows)

    # Summary
    for row in import_rows:
        if row.is_valid:
            batch.valid_rows += 1
        else:
            batch.failed_rows += 1
        if row.calculated and row.result:
            batch.calculated_rows += 1
            if row.result.is_profitable:
                batch.profitable_rows += 1
            else:
                batch.loss_rows += 1

    batch.rows = import_rows
    return batch


def import_from_xlsx(file_path: str) -> ImportBatchResult:
    """Import and validate products from an XLSX file."""
    batch = ImportBatchResult(timestamp=datetime.now().isoformat(), filename=file_path)

    try:
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        ws = wb.active

        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            batch.errors.append("Spreadsheet is empty.")
            return batch

        headers = [str(h).lower().strip().replace(" ", "_") if h else "" for h in rows[0]]
        data_rows = rows[1:]
        
        # If the first data row appears to be the descriptive helper row (e.g. from generate_import_template), skip it
        if data_rows and any("description" in str(v).lower() or "cost" in str(v).lower() and not str(v).replace(".", "").isdigit() for v in data_rows[0]):
            first_val = str(data_rows[0][0]).lower() if data_rows[0] else ""
            if "product name" in first_val or "description" in first_val:
                data_rows = data_rows[1:]

        batch.total_rows = len(data_rows)

        import_rows = []
        for i, row_tuple in enumerate(data_rows, start=2):
            row_data = {headers[j]: str(v) if v is not None else ""
                        for j, v in enumerate(row_tuple) if j < len(headers)}
            import_row = validate_row(row_data, i)
            import_rows.append(import_row)

        import_rows = calculate_batch(import_rows)

        for row in import_rows:
            if row.is_valid:
                batch.valid_rows += 1
            else:
                batch.failed_rows += 1
            if row.calculated and row.result:
                batch.calculated_rows += 1
                if row.result.is_profitable:
                    batch.profitable_rows += 1
                else:
                    batch.loss_rows += 1

        batch.rows = import_rows
        wb.close()

    except Exception as e:
        batch.errors.append(f"Failed to read XLSX: {str(e)}")

    return batch


def export_to_csv(batch: ImportBatchResult, output_path: str) -> bool:
    """Export batch calculation results to CSV."""
    try:
        fieldnames = [
            "row", "product_name", "sku", "category", "channel",
            "purchase_cost", "actual_selling_price", "quantity",
            "net_revenue", "total_acquisition_cost", "total_marketplace_fees",
            "total_fulfillment_cost", "advertising_expense",
            "contribution_profit_before_adv", "contribution_profit_after_adv",
            "estimated_operating_profit", "contribution_margin_pct",
            "operating_margin_pct", "markup_on_cost_pct",
            "is_profitable", "status", "errors", "warnings"
        ]
        with open(output_path, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
            writer.writeheader()
            for row in batch.rows:
                record = {
                    "row": row.row_number,
                    "product_name": row.parsed_data.get("product_name", ""),
                    "sku": row.parsed_data.get("sku", ""),
                    "category": row.parsed_data.get("category", ""),
                    "channel": row.parsed_data.get("channel", ""),
                    "purchase_cost": row.parsed_data.get("purchase_cost", ""),
                    "actual_selling_price": row.parsed_data.get("actual_selling_price", ""),
                    "quantity": row.parsed_data.get("quantity", 1),
                    "errors": "; ".join(row.errors),
                    "warnings": "; ".join(row.warnings),
                }
                if row.result:
                    r = row.result
                    record.update({
                        "net_revenue": str(r.net_revenue),
                        "total_acquisition_cost": str(r.total_acquisition_cost),
                        "total_marketplace_fees": str(r.total_marketplace_fees),
                        "total_fulfillment_cost": str(r.total_fulfillment_cost),
                        "advertising_expense": str(r.advertising_expense),
                        "contribution_profit_before_adv": str(r.contribution_profit_before_adv),
                        "contribution_profit_after_adv": str(r.contribution_profit_after_adv),
                        "estimated_operating_profit": str(r.estimated_operating_profit),
                        "contribution_margin_pct": str(r.contribution_margin_pct),
                        "operating_margin_pct": str(r.operating_margin_pct),
                        "markup_on_cost_pct": str(r.markup_on_cost_pct),
                        "is_profitable": str(r.is_profitable),
                        "status": "PROFIT" if r.is_profitable else "LOSS",
                    })
                else:
                    record["status"] = "ERROR" if not row.is_valid else "NOT_CALCULATED"
                writer.writerow(record)
        return True
    except Exception as e:
        return False


def generate_import_template(output_path: str) -> bool:
    """Generate a sample import template with example data."""
    try:
        wb = Workbook()
        ws = wb.active
        ws.title = "Products"

        # Headers
        headers = list(COLUMN_SPEC.keys())
        ws.append(headers)

        # Descriptions row
        descriptions = [COLUMN_SPEC[h]["description"] for h in headers]
        ws.append(descriptions)

        # Sample rows (clearly marked as examples)
        samples = [
            {
                "product_name": "SAMPLE: Cotton T-Shirt (M)",
                "sku": "TSH-001",
                "purchase_cost": "280",
                "packaging_cost": "15",
                "inbound_shipping": "20",
                "listed_price": "799",
                "actual_selling_price": "699",
                "quantity": "10",
                "marketplace_commission": "49",
                "forward_shipping": "65",
                "payment_gateway_fee": "14",
                "advertising_cost": "30",
                "gst_rate_pct": "5",
                "category": "Apparel",
                "channel": "Amazon India",
            },
            {
                "product_name": "SAMPLE: Bluetooth Earbuds",
                "sku": "EAR-002",
                "purchase_cost": "550",
                "packaging_cost": "30",
                "inbound_shipping": "40",
                "listed_price": "1499",
                "actual_selling_price": "1299",
                "quantity": "5",
                "marketplace_commission": "130",
                "forward_shipping": "85",
                "payment_gateway_fee": "26",
                "advertising_cost": "80",
                "gst_rate_pct": "18",
                "category": "Electronics",
                "channel": "Flipkart",
            },
        ]

        for sample in samples:
            ws.append([sample.get(h, "") for h in headers])

        # Style header row
        from openpyxl.styles import Font, PatternFill
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="312E81")

        # Style description row
        for cell in ws[2]:
            cell.font = Font(italic=True, color="64748B")

        wb.save(output_path)
        return True
    except Exception as e:
        return False
